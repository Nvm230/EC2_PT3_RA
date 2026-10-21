#!/usr/bin/env python3
"""
===============================================================================
GENERADOR DE PRESENTACIÓN PPTX PROFESIONAL (5 A 8 MINUTOS) - ALTA CALIDAD
===============================================================================
Basado en la plantilla oficial de UTEC: /home/nvm/Downloads/Presentacion_MT5103_S16.pptx
Genera una presentación de 7 diapositivas con:
- Portada oficial de UTEC personalizada.
- Selección y justificación de Unitree A1 y Unitree H1 con retratos 3D renderizados.
- Modelado cinemático, formulación analítica de Jacobiano y benchmark numérico.
- Control cinemático cuadrúpedo con marcha Crawl quíntica y video MP4 embebido.
- Control humanoide con Meme '6-7' (brazos alternados + cintura) y agache profundo.
- Simulación temporal con gráficas de seguimiento y video MP4 embebido del H1.
- Conclusiones, enlaces al repositorio GitHub y espacio para preguntas.
"""

import os
import shutil
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

SRC_TEMPLATE = "/home/nvm/Downloads/Presentacion_MT5103_S16.pptx"
OUT_PPTX = "/home/nvm/Documents/UTEC/Avanzada_tareas/presentacion_robotica_avanzada.pptx"

# Copiar plantilla base
shutil.copyfile(SRC_TEMPLATE, OUT_PPTX)
prs = pptx.Presentation(OUT_PPTX)

# Paleta cromática corporativa y técnica
COLOR_TITLE = RGBColor(15, 23, 42)       # Slate 900 (Azul oscuro profundo)
COLOR_SUBTITLE = RGBColor(71, 85, 105)   # Slate 600
COLOR_PRIMARY = RGBColor(0, 102, 204)    # Azul institucional UTEC
COLOR_CYAN = RGBColor(14, 165, 233)      # Celeste tecnológico
COLOR_GREEN = RGBColor(22, 163, 74)      # Verde éxito
COLOR_CARD_BG = RGBColor(248, 250, 252)  # Slate 50
COLOR_CARD_ALT = RGBColor(241, 245, 249) # Slate 100
COLOR_BORDER = RGBColor(226, 232, 240)   # Slate 200
COLOR_TEXT = RGBColor(30, 41, 59)        # Slate 800
COLOR_MUTED = RGBColor(100, 116, 139)    # Slate 500
COLOR_WHITE = RGBColor(255, 255, 255)

def delete_slide(prs, index):
    rId = prs.slides._sldIdLst[index].rId
    prs.part.drop_rel(rId)
    del prs.slides._sldIdLst[index]

# -----------------------------------------------------------------------------
# 1. SLIDE 1: PORTADA OFICIAL UTEC
# -----------------------------------------------------------------------------
slide1 = prs.slides[0]
for s in slide1.shapes:
    if s.name == "TextBox 6":
        tf = s.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = "Alegría, Marco Antonio"
        p.font.bold = True
        p.font.size = Pt(22)
        p.font.color.rgb = COLOR_TITLE
    elif s.name == "TextBox 11":
        s.text_frame.text = "Alumno:"
    elif s.name == "TextBox 12":
        s.text_frame.text = "Docente: Robótica Avanzada"
    elif s.name == "TextBox 13":
        s.text_frame.text = "Curso: Robótica Avanzada — Evaluación Continua 2 (Tarea 3)"
    elif s.name == "Group 7":
        for child in s.shapes:
            if child.name == "TextBox 8":
                child.text_frame.text = "2026-2"
            elif child.name == "TextBox 9":
                tf = child.text_frame
                tf.clear()
                p = tf.paragraphs[0]
                p.text = "CINEMÁTICA Y CONTROL DE ROBOTS CUADRÚPEDO (UNITREE A1) Y HUMANOIDE (UNITREE H1) CON PROXQP"
                p.font.bold = True
                p.font.size = Pt(28)
                p.font.color.rgb = COLOR_TITLE

# -----------------------------------------------------------------------------
# UTILIDADES DE LAYOUT
# -----------------------------------------------------------------------------
def clean_slide_content(slide):
    """Conserva únicamente el fondo blanco (0), el footer inferior (1) y el logo UTEC (2)."""
    to_remove = [s for s in list(slide.shapes)[3:]]
    for s in to_remove:
        sp = s._element
        sp.getparent().remove(sp)

def add_header(slide, title_text, subtitle_text, slide_num):
    """Genera cabecera estandarizada sin superposiciones."""
    # Título principal
    txBox = slide.shapes.add_textbox(Inches(1.1), Inches(0.55), Inches(15.5), Inches(0.75))
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.bold = True
    p.font.size = Pt(26)
    p.font.color.rgb = COLOR_TITLE
    
    # Subtítulo explicativo
    txBox2 = slide.shapes.add_textbox(Inches(1.1), Inches(1.35), Inches(15.5), Inches(0.45))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
    p2 = tf2.paragraphs[0]
    p2.text = subtitle_text
    p2.font.size = Pt(15)
    p2.font.color.rgb = COLOR_SUBTITLE

    # Número de página
    txBox3 = slide.shapes.add_textbox(Inches(18.2), Inches(10.0), Inches(1.0), Inches(0.4))
    tf3 = txBox3.text_frame
    tf3.margin_left = tf3.margin_top = tf3.margin_right = tf3.margin_bottom = 0
    p3 = tf3.paragraphs[0]
    p3.text = str(slide_num)
    p3.alignment = PP_ALIGN.RIGHT
    p3.font.size = Pt(14)
    p3.font.color.rgb = COLOR_MUTED

def add_card(slide, left, top, width, height, title, items, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER, title_color=COLOR_TITLE):
    """Crea una tarjeta contenedora estilizada con sombra y viñetas cuidadas."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.22)
    tf.margin_right = Inches(0.25)
    tf.margin_bottom = Inches(0.2)
    tf.word_wrap = True
    tf.clear()
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.bold = True
    p.font.size = Pt(17)
    p.font.color.rgb = title_color
    p.space_after = Pt(8)
    
    for item in items:
        p2 = tf.add_paragraph()
        p2.text = f"• {item}"
        p2.font.size = Pt(13.5)
        p2.font.color.rgb = COLOR_TEXT
        p2.space_after = Pt(5)
        
    return shape

def add_badge_card(slide, left, top, width, height, title, metrics):
    """Crea una tarjeta de métricas clave compacta."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_CARD_ALT
    shape.line.color.rgb = COLOR_BORDER
    shape.line.width = Pt(1.0)
    
    tf = shape.text_frame
    tf.margin_left = Inches(0.2)
    tf.margin_top = Inches(0.18)
    tf.margin_right = Inches(0.2)
    tf.margin_bottom = Inches(0.15)
    tf.word_wrap = True
    tf.clear()
    
    p = tf.paragraphs[0]
    p.text = title.upper()
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(6)
    
    for k, v in metrics:
        p2 = tf.add_paragraph()
        p2.text = f"{k}: "
        p2.font.bold = True
        p2.font.size = Pt(12.5)
        p2.font.color.rgb = COLOR_TITLE
        run = p2.add_run()
        run.text = v
        run.font.bold = False
        run.font.color.rgb = COLOR_TEXT
        p2.space_after = Pt(4)
    return shape

def add_video_container(slide, left, top, width, height, video_file, poster_file, title_text, status_text):
    """Inserta reproductor de video nativo con póster y cintillo descriptivo."""
    # Intentar embeber video con poster
    if os.path.exists(video_file):
        try:
            slide.shapes.add_movie(
                video_file,
                left, top, width, height,
                poster_frame_image=poster_file if (poster_file and os.path.exists(poster_file)) else None,
                mime_type='video/mp4'
            )
        except Exception as e:
            print(f"Error embebiendo movie {video_file}: {e}")
            if poster_file and os.path.exists(poster_file):
                slide.shapes.add_picture(poster_file, left, top, width=width, height=height)
    elif poster_file and os.path.exists(poster_file):
        slide.shapes.add_picture(poster_file, left, top, width=width, height=height)

    # Cintillo superpuesto/inferior
    label_h = Inches(0.65)
    lbl = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top + height - label_h, width, label_h)
    lbl.fill.solid()
    lbl.fill.fore_color.rgb = RGBColor(15, 23, 42)
    lbl.line.fill.background()
    tf = lbl.text_frame
    tf.margin_left = Inches(0.15)
    tf.margin_top = Inches(0.08)
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"▶ {title_text}"
    p.font.bold = True
    p.font.size = Pt(12.5)
    p.font.color.rgb = RGBColor(56, 189, 248) # Cyan brillante
    p2 = tf.add_paragraph()
    p2.text = status_text
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = RGBColor(226, 232, 240)

# -----------------------------------------------------------------------------
# 2. SLIDE 2: PLATAFORMAS ROBÓTICAS Y JUSTIFICACIÓN
# -----------------------------------------------------------------------------
slide2 = prs.slides[1]
clean_slide_content(slide2)
add_header(slide2, "1. Plataformas Robóticas Seleccionadas y Justificación", 
           "Análisis comparativo de grados de libertad, arquitectura mecánica y objetivos de control", 2)

# Columna Izquierda: A1
if os.path.exists("outputs/unitree_a1_portrait.png"):
    slide2.shapes.add_picture("outputs/unitree_a1_portrait.png", Inches(1.1), Inches(2.1), width=Inches(4.2), height=Inches(4.0))

add_badge_card(slide2, Inches(1.1), Inches(6.3), Inches(4.2), Inches(3.6), 
               "Especificaciones A1", [
                   ("Grados de Libertad", "12 DOF (3 por pata)"),
                   ("Masa Total", "12.0 kg"),
                   ("Altura Nominal", "32.0 cm"),
                   ("Control Frecuencia", "100 Hz (dt = 0.01 s)"),
                   ("Cumplimiento", "Robot cuadrúpedo ≠ Laikago")
               ])

add_card(slide2, Inches(5.6), Inches(2.1), Inches(4.1), Inches(7.8), 
         "Unitree A1 (Cuadrúpedo)", [
             "Plataforma móvil cuadrúpeda de alta dinámica para locomoción terrestre.",
             "Justificación de selección: Seleccionado expresamente para cumplir con el requisito de utilizar un robot cuadrúpedo alternativo al modelo Laikago.",
             "Arquitectura mecánica: Cuatro cadenas seriales idénticas conectadas a una base flotante.",
             "Distribución articular: Cadera roll (abducción/aducción), cadera pitch (flexión) y rodilla pitch.",
             "Objetivo de la tarea: Demostración de cinemática analítica cerrada pura (NumPy) y marcha secuencial estática Crawl."
         ])

# Columna Derecha: H1
if os.path.exists("outputs/unitree_h1_portrait.png"):
    slide2.shapes.add_picture("outputs/unitree_h1_portrait.png", Inches(10.1), Inches(2.1), width=Inches(4.2), height=Inches(4.0))

add_badge_card(slide2, Inches(10.1), Inches(6.3), Inches(4.2), Inches(3.6), 
               "Especificaciones H1", [
                   ("Grados de Libertad", "19 DOF articulados"),
                   ("Estatura / Masa", "1.80 m / 47.0 kg"),
                   ("Cintura (Torso)", "1 DOF (yaw Z: ±25.8°)"),
                   ("Extremidades", "10 DOF piernas, 8 DOF brazos"),
                   ("Simulador", "PyBullet con visualización directa")
               ])

add_card(slide2, Inches(14.6), Inches(2.1), Inches(4.3), Inches(7.8), 
         "Unitree H1 (Humanoide)", [
             "Humanoide bípedo de escala real para investigación en manipulación móvil y locomoción bípeda.",
             "Justificación de selección: Robot de alta dimensionalidad (19 DOF) idóneo para optimización cuadrática multitarea.",
             "Análisis anatómico: Al carecer de articulación en cuello/cabeza, se utiliza el actuador torso_joint (yaw Z) para orientar el campo visual.",
             "Objetivo Movimiento A: Coordinación rítmica de brazos con giro oscilatorio de cintura (Meme '6-7').",
             "Objetivo Movimiento B: Agache profundo descendiendo pelvis 52 cm para recolección en piso."
         ])

# -----------------------------------------------------------------------------
# 3. SLIDE 3: MODELADO CINEMÁTICO A1 (FK Y JACOBIANO)
# -----------------------------------------------------------------------------
slide3 = prs.slides[2]
clean_slide_content(slide3)
add_header(slide3, "2. Modelado Cinemático y Jacobiano Analítico (Unitree A1)", 
           "Derivación matemática cerrada desde primeros principios (NumPy) y validación de precisión de máquina", 3)

add_card(slide3, Inches(1.1), Inches(2.1), Inches(8.6), Inches(3.6), 
         "Derivación Matemática Analítica (Primeros Principios)", [
             "Cinemática Directa (FK): Matrices homogéneas exactas Rx(q1) * Ry(q2) * Ry(q3) considerando offsets de cadera (dx=0.1805m, dy=0.047m) y eslabones (0.2m, 0.2m).",
             "Jacobiano Geométrico Analítico (12x12): Derivado analíticamente mediante producto vectorial Ji = zi x (p_pie - pi) para cada eje de rotación.",
             "Cero librerías externas: Implementación 100% matemática en NumPy puro, sin diferencias finitas ni librerías cinemáticas intermedias."
         ])

if os.path.exists("outputs/part1_estructura_robot_3d.png"):
    slide3.shapes.add_picture("outputs/part1_estructura_robot_3d.png", Inches(1.1), Inches(5.9), width=Inches(8.6), height=Inches(4.0))

add_card(slide3, Inches(10.1), Inches(2.1), Inches(8.8), Inches(3.6), 
         "Benchmark Numérico vs Pinocchio (200 Muestras)", [
             "200 configuraciones articulares aleatorias uniformes sobre el espacio de trabajo físico completo.",
             "Error cartesiano máximo en FK: 1.11 x 10^-16 metros (límite físico IEEE 754 de doble precisión float64).",
             "Error máximo en Jacobiano: 1.11 x 10^-16 s^-1, demostrando identidad algebraica perfecta.",
             "Velocidad de cómputo analítico: < 45 microsegundos por configuración, apto para control en tiempo real."
         ])

if os.path.exists("outputs/part1_errores_benchmark.png"):
    slide3.shapes.add_picture("outputs/part1_errores_benchmark.png", Inches(10.1), Inches(5.9), width=Inches(8.8), height=Inches(4.0))

# -----------------------------------------------------------------------------
# 4. SLIDE 4: CONTROL CINEMÁTICO Y LOCOMOCIÓN DEL A1 (CON VIDEO)
# -----------------------------------------------------------------------------
slide4 = prs.slides[3]
clean_slide_content(slide4)
add_header(slide4, "3. Control Cinemático y Locomoción Cuadrúpeda (Unitree A1)", 
           "Inversión cinemática en lazo cerrado (CLIK) con ProxQP a 100 Hz y marcha secuencial estática", 4)

add_card(slide4, Inches(1.1), Inches(2.1), Inches(8.6), Inches(4.1), 
         "Control por Optimización Cuadrática (ProxQP a 100 Hz)", [
             "Ley CLIK en lazo cerrado: v_ref = v_des + Kp * (p_des - p_actual), con ganancia proporcional Kp = 20.0.",
             "Formulación QP:\n  min  0.5 || J dq - v_ref ||^2_W + 0.5 lambda ||dq||^2\n  s.a. dq_min <= dq <= dq_max,  (q_min - q)/dt <= dq <= (q_max - q)/dt",
             "Marcha secuencial Crawl: FL -> RR -> FR -> RL, garantizando siempre soporte estático en trípode.",
             "Perfil quíntico de balanceo (Swing): Elevación vertical h = 8.0 cm y avance horizontal Delta x = 10.0 cm.",
             "Precisión cartesiana submilimétrica: Error medio de 0.54 mm y máximo de 0.89 mm en todo el ciclo."
         ])

add_video_container(slide4, Inches(1.1), Inches(6.4), Inches(8.6), Inches(3.5),
                    "outputs/part2_a1_gait.mp4", "outputs/part2_step_2.png",
                    "Video: Locomoción Cuadrúpeda Unitree A1",
                    "Marcha Crawl 4 pasos | Cámara dinámica y HUD de soporte")

if os.path.exists("outputs/part2_all_steps_grid.png"):
    slide4.shapes.add_picture("outputs/part2_all_steps_grid.png", Inches(10.1), Inches(2.1), width=Inches(8.8), height=Inches(7.8))

# -----------------------------------------------------------------------------
# 5. SLIDE 5: CONTROL DEL HUMANOIDE H1 (MOVIMIENTOS A Y B)
# -----------------------------------------------------------------------------
slide5 = prs.slides[4]
clean_slide_content(slide5)
add_header(slide5, "4. Modelado y Control Cinemático de Humanoide (Unitree H1)", 
           "Control de 19 DOF con tareas prioritarias de manipulación cartesiana, giro de cintura y agache", 5)

# Movimiento A (Izquierda)
add_card(slide5, Inches(1.1), Inches(2.1), Inches(8.6), Inches(4.0), 
         "Movimiento A: Brazos Alternados '6-7' y Cintura Yaw", [
             "Articulación de cintura: El H1 posee torso_joint (yaw Z, articulación 10). Se controla explícitamente para orientar el cuerpo.",
             "Tarea prioritaria en QP (peso 2.5): Oscilación sinusoidal amplia de +/- 25.8° (+/- 0.45 rad) a f = 0.5 Hz.",
             "Gesto rítmico '6-7':\n  - Semiciclo derecho: Brazo derecho sube +35 cm mientras torso gira a la derecha.\n  - Semiciclo izquierdo: Brazo izquierdo sube +35 cm mientras torso gira a la izquierda.",
             "Coordinación natural: La cinemática distribuye el movimiento entre hombros y codos sin singularidades."
         ])

if os.path.exists("outputs/part3_dual_arms_waist.png"):
    slide5.shapes.add_picture("outputs/part3_dual_arms_waist.png", Inches(1.1), Inches(6.3), width=Inches(8.6), height=Inches(3.6))

# Movimiento B (Derecha)
add_card(slide5, Inches(10.1), Inches(2.1), Inches(8.8), Inches(4.0), 
         "Movimiento B: Agache Profundo y Recolección en Piso", [
             "Objetivo físico: Alcanzar cilindro en el suelo en [0.28, -0.18, 0.08] m (altura objetivo z = 0.16 m).",
             "Descenso guiado de pelvis: Reducción vertical de 52 cm (desde z=0.98m hasta z=0.46m).",
             "Estabilidad del Centro de Masas (CoM): Inclinación frontal coordinada del torso (pitch = 31.5°) para mantener la proyección del CoM dentro del polígono de soporte plantar.",
             "Cinemática de piernas: Flexión articular simétrica de rodillas (111.7°), caderas (-88.8°) y tobillos (-54.4°).",
             "Retorno seguro: Reincorporación completa a postura erguida con error nulo."
         ])

if os.path.exists("outputs/part3_crouch_floor.png"):
    slide5.shapes.add_picture("outputs/part3_crouch_floor.png", Inches(10.1), Inches(6.3), width=Inches(8.8), height=Inches(3.6))

# -----------------------------------------------------------------------------
# 6. SLIDE 6: SIMULACIÓN HUMANOIDE H1 (CON VIDEO Y GRÁFICAS)
# -----------------------------------------------------------------------------
slide6 = prs.slides[5]
clean_slide_content(slide6)
add_header(slide6, "5. Simulación del Humanoide H1 y Desempeño Temporal", 
           "Validación en PyBullet y curvas temporales de seguimiento de pelvis, mano derecha y cintura", 6)

if os.path.exists("outputs/part3_motion_sequence_grid.png"):
    slide6.shapes.add_picture("outputs/part3_motion_sequence_grid.png", Inches(1.1), Inches(2.1), width=Inches(8.6), height=Inches(4.4))

add_video_container(slide6, Inches(1.1), Inches(6.7), Inches(8.6), Inches(3.2),
                    "outputs/part3_h1_humanoid.mp4", "outputs/part3_standing.png",
                    "Video: Demostración Completa Humanoide H1",
                    "Postura erguida -> Gesto 6-7 con cintura -> Agache al piso -> Retorno")

if os.path.exists("outputs/part3_h1_tracking.png"):
    slide6.shapes.add_picture("outputs/part3_h1_tracking.png", Inches(10.1), Inches(2.1), width=Inches(8.8), height=Inches(7.8))

# -----------------------------------------------------------------------------
# 7. SLIDE 7: CONCLUSIONES Y PREGUNTAS
# -----------------------------------------------------------------------------
slide7 = prs.slides[13]

# Eliminar textbox antiguo de gracias para rediseñarlo
for s in list(slide7.shapes):
    if s.name == "TextBox 4" or (s.has_text_frame and "Gracias" in s.text):
        sp = s._element
        sp.getparent().remove(sp)

# Título de cierre
tx_close = slide7.shapes.add_textbox(Inches(1.1), Inches(0.8), Inches(11.0), Inches(1.5))
tf_c = tx_close.text_frame
tf_c.word_wrap = True
tf_c.margin_left = tf_c.margin_top = 0
p_c1 = tf_c.paragraphs[0]
p_c1.text = "¡Muchas Gracias!"
p_c1.font.bold = True
p_c1.font.size = Pt(36)
p_c1.font.color.rgb = COLOR_TITLE
p_c2 = tf_c.add_paragraph()
p_c2.text = "Espacio para Preguntas y Discusión Técnica"
p_c2.font.size = Pt(20)
p_c2.font.color.rgb = COLOR_PRIMARY

# Tarjeta de Conclusiones Clave
add_card(slide7, Inches(1.1), Inches(2.6), Inches(9.8), Inches(4.3),
         "Conclusiones y Logros Clave", [
             "Modelado Analítico Exacto: Formulación analítica en NumPy con error 1.11 x 10^-16 m (límite numérico IEEE 754), validando la cinemática de extremidades cuadrúpedas.",
             "Control Cinemático Robusto: Formulación cuadrática ProxQP a 100 Hz con manejo estricto de límites de velocidad y posición articular.",
             "Locomoción Estable: Marcha Crawl quíntica con soporte en trípode continuo y elevación vertical limpia de 8 cm.",
             "Humanoide de Cuerpo Completo: Solución coordinada multitarea para 19 DOF integrando giro de cintura yaw (±25.8°) y agache profundo para recolección en piso."
         ], bg_color=COLOR_WHITE)

# Tarjeta de Repositorio GitHub y Autor
repo_card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), Inches(7.1), Inches(9.8), Inches(1.5))
repo_card.fill.solid()
repo_card.fill.fore_color.rgb = RGBColor(15, 23, 42)
repo_card.line.color.rgb = RGBColor(0, 102, 204)
repo_card.line.width = Pt(1.5)
tf_r = repo_card.text_frame
tf_r.margin_left = Inches(0.2)
tf_r.margin_top = Inches(0.15)
tf_r.word_wrap = True
p_r1 = tf_r.paragraphs[0]
p_r1.text = "REPOSITORIO DEL PROYECTO (CÓDIGO, VIDEOS Y REPORTES)"
p_r1.font.bold = True
p_r1.font.size = Pt(12)
p_r1.font.color.rgb = RGBColor(56, 189, 248)
p_r2 = tf_r.add_paragraph()
p_r2.text = "GitHub: https://github.com/Nvm230/EC2_PT3_RA"
p_r2.font.bold = True
p_r2.font.size = Pt(14)
p_r2.font.color.rgb = RGBColor(255, 255, 255)
p_r3 = tf_r.add_paragraph()
p_r3.text = "Autor: Marco Antonio Alegría | UTEC — Octubre 2026"
p_r3.font.size = Pt(12)
p_r3.font.color.rgb = RGBColor(203, 213, 225)

# -----------------------------------------------------------------------------
# ELIMINAR DIAPOSITIVAS SOBRANTES DE LA PLANTILLA
# -----------------------------------------------------------------------------
slides_to_delete = [16, 15, 14, 12, 11, 10, 9, 8, 7, 6]
for idx in slides_to_delete:
    delete_slide(prs, idx)

prs.save(OUT_PPTX)
print(f"Presentación generada exitosamente en: {OUT_PPTX}")
print(f"Total diapositivas finales: {len(prs.slides)}")
